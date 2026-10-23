#!/usr/bin/env python3
"""Rename files and folders using an Excel mapping. See README.md for usage."""
import argparse
import hashlib
import json
import os
import sys
import threading
import uuid
from collections import Counter
from pathlib import Path

STATE_NAME = '.rename-state.json'
LOCK_NAME = '.rename.lock'
PROJECT_ROOT = Path(__file__).resolve().parent.parent


def validate_basename(value):
    """Use portable basenames, including Windows filename restrictions."""
    reserved = {'CON', 'PRN', 'AUX', 'NUL', 'CONIN$', 'CONOUT$'}
    reserved.update(f'{prefix}{n}' for prefix in ('COM', 'LPT') for n in '123456789\u00b9\u00b2\u00b3')
    if (not value or value in ('.', '..') or value.endswith((' ', '.'))
            or any(ord(c) < 32 or c in '<>:"/\\|?*' for c in value)
            or value.split('.')[0].upper() in reserved
            or value.casefold() in (STATE_NAME, LOCK_NAME)
            or value.startswith('.__rename_')):
        raise ValueError(f'Unsafe basename: {value!r}')


def validate_config(config):
    if not isinstance(config, dict):
        raise ValueError('Configuration must be a JSON object')
    for key in ('root_directory', 'excel_path'):
        if not isinstance(config.get(key), str) or not config[key].strip():
            raise ValueError(f'{key} must be a nonempty path string')
    for key in ('dry_run', 'recursive', 'rename_files', 'rename_folders'):
        if key in config and type(config[key]) is not bool:
            raise ValueError(f'{key} must be true or false (a JSON boolean)')
    for key in ('source_column', 'destination_column'):
        if key in config and (not isinstance(config[key], str) or not config[key].strip()):
            raise ValueError(f'{key} must be a nonempty column name')
    if 'history_directory' in config and (not isinstance(config['history_directory'], str)
                                         or not config['history_directory'].strip()):
        raise ValueError('history_directory must be a nonempty path string')
    if config.get('sheet_name') is not None and not isinstance(config['sheet_name'], str):
        raise ValueError('sheet_name must be a string or null')
    extensions = config.get('file_extensions')
    if (not isinstance(extensions, list) or not extensions
            or any(not isinstance(ext, str) or not ext.startswith('.') or len(ext) < 2
                   or any(ord(c) < 33 or c in '<>:"/\\|?*.' for c in ext[1:])
                   for ext in extensions)):
        raise ValueError('file_extensions must list extensions like .pdf')


def load_mapping(excel_path, sheet_name, source_column='New', destination_column='Equivalent from old'):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("Install dependency: python -m pip install openpyxl")
    book = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in book.sheetnames:
            raise ValueError(f'Mapping sheet not found: {sheet_name}')
        sheet = book[sheet_name] if sheet_name else book.active
        rows = iter(sheet.values)
        headers = next(rows, None)
        if not headers:
            raise ValueError('Mapping sheet is empty')
        headers = [str(x).strip().casefold() if x is not None else '' for x in headers]
        try:
            ni, oi = headers.index(source_column.strip().casefold()), headers.index(destination_column.strip().casefold())
        except ValueError as exc:
            raise ValueError(f'Expected columns "{source_column}" and "{destination_column}"') from exc
        mapping = {}
        for line, row in enumerate(rows, 2):
            new = str(row[ni]).strip() if len(row) > ni and row[ni] is not None else ''
            old = str(row[oi]).strip() if len(row) > oi and row[oi] is not None else ''
            if not new and not old:
                continue
            if not new or not old:
                raise ValueError(f'Incomplete row {line}')
            for value in (new, old):
                try:
                    validate_basename(value)
                except ValueError as exc:
                    raise ValueError(f'{exc} at row {line}') from exc
            if new in mapping and mapping[new] != old:
                raise ValueError(f'Conflicting mapping for {new} at row {line}')
            mapping[new] = old
        if not mapping:
            raise ValueError('Excel mapping is empty')
        return mapping
    finally:
        book.close()


def target_name(name, mapping, file, extensions):
    if file:
        extension = Path(name).suffix
        if extension.lower() not in extensions:
            return name
        stem = name[:-len(extension)]
        return mapping.get(stem, stem) + extension
    return mapping.get(name, name)


def identity(path):
    stat = path.stat(follow_symlinks=False)
    if not stat.st_ino:
        raise ValueError(f'Filesystem does not provide stable file identities: {path}')
    return f'{stat.st_dev}:{stat.st_ino}:{getattr(stat, "st_birthtime_ns", 0)}'


def plan_group(folder, names, mapping, file, extensions, counts, processed=(),
               existing=None, identities=None):
    planned = []
    for name in names:
        original = folder / name
        if name in (STATE_NAME, LOCK_NAME):
            continue
        if name.startswith('.__rename_'):
            raise ValueError(f'Unrecovered temporary rename found: {original}')
        destination = target_name(name, mapping, file, extensions)
        if destination == name:
            continue
        # Unmatched names need no filesystem or history checks.
        if (original.is_symlink()
                or getattr(original, 'is_junction', lambda: False)()):
            continue
        validate_basename(destination)
        if processed or identities is not None:
            source_id = identity(original)
            if source_id in processed:
                continue
            if identities is not None:
                identities[name] = source_id
        planned.append((name, destination))
    if not planned:
        return []

    if existing is None:
        existing = set(os.listdir(folder))
    sources = {old for old, _ in planned}
    # Windows is case-insensitive by default; Linux may be case-sensitive.
    normalize = str.casefold if os.name == 'nt' else lambda x: x
    targets = Counter(normalize(dst) for _, dst in planned)
    occupied = {normalize(name) for name in existing if name not in sources}
    conflicts = [(src, dst) for src, dst in planned
                 if targets[normalize(dst)] > 1 or normalize(dst) in occupied]
    if conflicts:
        counts['conflicts'] += len(conflicts)
        for src, dst in conflicts:
            print(f'CONFLICT: {folder / src} -> {folder / dst}')
        print(f'SKIP group: {folder} (no changes in this group)')
        return []

    return planned


def temporary_name(folder):
    while True:
        name = f'.__rename_{uuid.uuid4().hex}'
        if not os.path.lexists(folder / name):
            return name


def move(source, destination):
    # Path.rename can overwrite on POSIX; refuse occupied destinations everywhere.
    if os.path.lexists(destination):
        raise FileExistsError(f'Destination already exists: {destination}')
    source.rename(destination)


def restore_group(folder, entries, locations):
    """Free all original names before restoring, including swaps and cycles."""
    for index, (src, temp, dst) in enumerate(entries):
        if locations[index] == dst:
            move(folder / dst, folder / temp)
            locations[index] = temp
    for index, (src, temp, dst) in enumerate(entries):
        if locations[index] == temp:
            move(folder / temp, folder / src)
            locations[index] = src


def execute_group(folder, entries):
    locations = [src for src, _, _ in entries]
    try:
        for index, (src, temp, dst) in enumerate(entries):
            move(folder / src, folder / temp)
            locations[index] = temp
        for index, (src, temp, dst) in enumerate(entries):
            move(folder / temp, folder / dst)
            locations[index] = dst
    except OSError:
        try:
            restore_group(folder, entries, locations)
        except OSError as exc:
            raise RuntimeError(f'RECOVERY NEEDED in {folder}: {exc}') from exc
        raise


def rename_group(folder, names, mapping, file, extensions, apply, counts):
    """Single-group API; command-line runs also persist history and undo earlier groups."""
    planned = plan_group(folder, names, mapping, file, extensions, counts)

    for src, dst in planned:
        print(f'{"APPLY" if apply else "PREVIEW"}: {folder / src} -> {folder / dst}')
    counts['planned'] += len(planned)
    if not apply:
        return

    entries = [(src, temporary_name(folder), dst) for src, dst in planned]
    try:
        execute_group(folder, entries)
        counts['renamed'] += len(entries)
    except (OSError, RuntimeError) as exc:
        counts['errors'] += 1
        print(f'ERROR: {exc}', file=sys.stderr)


def history_paths(root, history_directory):
    history_directory = history_directory.resolve()
    if history_directory.is_relative_to(root):
        raise ValueError('history_directory must be outside root_directory')
    key = hashlib.sha256(os.path.normcase(str(root)).encode()).hexdigest()
    return history_directory / f'{key}.json', history_directory / f'{key}.lock'


def read_state(root, path):
    if not path.exists():
        return {'version': 1, 'root': str(root), 'processed': {}, 'pending': []}
    state = json.loads(path.read_text(encoding='utf-8'))
    if (not isinstance(state, dict) or state.get('version') != 1
            or state.get('root') != str(root) or not isinstance(state.get('processed'), dict)
            or not isinstance(state.get('pending'), list)
            or any(not isinstance(v, list) or any(not isinstance(i, str) for i in v)
                   for v in state['processed'].values())):
        raise ValueError(f'Invalid rename history: {path}')
    if state['pending']:
        raise ValueError(f'RECOVERY NEEDED: interrupted run recorded in {path}; inspect its pending moves before continuing')
    return state


def write_state(state_path, state):
    path = state_path.parent / temporary_name(state_path.parent)
    try:
        with path.open('x', encoding='utf-8') as handle:
            json.dump(state, handle, indent=2)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(path, state_path)
    finally:
        if path.exists():
            path.unlink()


def run(root, config, mapping, apply, state_path):
    def walk_error(exc):
        raise exc

    # Finish traversal and collision checks before applying any rename.
    if config.get('recursive', True):
        walker = []
        for location, folders, files in os.walk(root, topdown=True, followlinks=False, onerror=walk_error):
            folder = Path(location)
            # Include skipped links/junctions in collision checks as occupied names.
            existing = set(folders) | set(files)
            folders[:] = [name for name in folders if not (folder / name).is_symlink()
                          and not getattr(folder / name, 'is_junction', lambda: False)()]
            walker.append((location, list(folders), files, existing))
        walker.reverse()
    else:
        with os.scandir(root) as scan:
            children = list(scan)
        walker = [(str(root), [p.name for p in children if p.is_dir(follow_symlinks=False)],
                   [p.name for p in children if p.is_file(follow_symlinks=False)],
                   {p.name for p in children})]
    state = read_state(root, state_path)
    fingerprint = hashlib.sha256(json.dumps(sorted(mapping.items()), ensure_ascii=True).encode()).hexdigest()
    processed = set(state['processed'].get(fingerprint, []))
    extensions = {ext.lower() for ext in config['file_extensions']}
    counts = dict(planned=0, renamed=0, conflicts=0, errors=0)
    groups = []
    new_ids = set()
    for location, folders, files, existing in walker:
        folder = Path(location)
        identities = {}
        for enabled, names, file in [(config.get('rename_files', True), files, True),
                                      (config.get('rename_folders', True), folders, False)]:
            if not enabled:
                continue
            planned = plan_group(folder, names, mapping, file, extensions, counts, processed,
                                 existing=existing, identities=identities)
            if planned:
                entries = [(src, temporary_name(folder), dst) for src, dst in planned]
                groups.append((folder, entries))
                new_ids.update(identities[src] for src, _ in planned)
                for src, dst in planned:
                    print(f'{"APPLY" if apply else "PREVIEW"}: {folder / src} -> {folder / dst}')
                counts['planned'] += len(planned)
    if apply and groups and not counts['conflicts']:
        state['pending'] = [{'folder': str(folder.relative_to(root)), 'moves': entries}
                            for folder, entries in groups]
        write_state(state_path, state)
        completed = []
        try:
            for folder, entries in groups:
                execute_group(folder, entries)
                completed.append((folder, entries))
            state['processed'][fingerprint] = sorted(processed | new_ids)
            state['pending'] = []
            write_state(state_path, state)
            counts['renamed'] = counts['planned']
        except OSError as exc:
            counts['errors'] += 1
            print(f'ERROR: {exc}; rolling back this run', file=sys.stderr)
            for folder, entries in reversed(completed):
                try:
                    restore_group(folder, entries, [dst for _, _, dst in entries])
                except OSError as restore_error:
                    raise RuntimeError(f'RECOVERY NEEDED: inspect {state_path}: {restore_error}') from restore_error
            state['processed'][fingerprint] = sorted(processed)
            state['pending'] = []
            write_state(state_path, state)
    print('Summary: ' + ', '.join(f'{k}={v}' for k, v in counts.items()))
    return 1 if counts['conflicts'] or counts['errors'] else 0


def _main():
    parser = argparse.ArgumentParser(description='Rename files and folders using an Excel mapping')
    parser.add_argument('--config', type=Path, default=PROJECT_ROOT / 'config' / 'config.json')
    parser.add_argument('--apply', action='store_true', help='Override dry_run and perform renaming')
    parser.add_argument('--dry-run', action='store_true', help='Preview even if dry_run is false in config')
    args = parser.parse_args()
    try:
        config_file = args.config.expanduser().resolve()
        if args.apply and args.dry_run:
            raise ValueError('Choose either --apply or --dry-run')
        config = json.loads(config_file.read_text(encoding='utf-8-sig'))
        validate_config(config)
        base = config_file.parent
        root = (base / config['root_directory']).resolve()
        excel = (base / config['excel_path']).resolve()
        if not root.is_dir():
            raise ValueError(f'Root directory not found: {root}')
        if not excel.is_file():
            raise ValueError(f'Excel mapping not found: {excel}')
        mapping = load_mapping(excel, config.get('sheet_name'), config.get('source_column', 'New'),
                               config.get('destination_column', 'Equivalent from old'))
        apply = not args.dry_run and (args.apply or not config.get('dry_run', True))
        print(f'Mode: {"APPLY" if apply else "DRY RUN"} | Root: {root} | Mappings: {len(mapping)}')
        history_directory = ((base / config['history_directory']) if 'history_directory' in config
                             else PROJECT_ROOT / '.rename-history')
        state_path, lock = history_paths(root, history_directory)
        legacy = root / STATE_NAME
        if os.path.lexists(lock) or os.path.lexists(root / LOCK_NAME):
            raise ValueError(f'Run lock exists: {lock}; check for another or interrupted run')
        if legacy.exists() and state_path.exists():
            raise ValueError(f'Two history files found: {legacy} and {state_path}; reconcile them before continuing')
        if apply:
            state_path.parent.mkdir(parents=True, exist_ok=True)
            with lock.open('x', encoding='utf-8') as handle:
                handle.write(str(os.getpid()))
            try:
                if legacy.exists():
                    read_state(root, legacy)
                    move(legacy, state_path)
                return run(root, config, mapping, apply, state_path)
            finally:
                lock.unlink()
        return run(root, config, mapping, apply, legacy if legacy.exists() else state_path)
    except (ValueError, KeyError, OSError, RuntimeError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 2


def main():
    print('Working... please wait.', file=sys.stderr, flush=True)
    finished = threading.Event()

    def report_activity():
        while not finished.wait(5):
            print('Still working... please wait.', file=sys.stderr, flush=True)

    reporter = threading.Thread(target=report_activity, daemon=True)
    reporter.start()
    try:
        return _main()
    finally:
        finished.set()
        reporter.join()


if __name__ == '__main__':
    sys.exit(main())
